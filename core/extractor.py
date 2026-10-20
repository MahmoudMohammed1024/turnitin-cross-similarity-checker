import os
import re
from typing import List, Dict, Any, Optional

try:
    import docx
except ImportError:
    docx = None

try:
    import pymupdf as fitz
except ImportError:
    try:
        import fitz
    except ImportError:
        fitz = None

try:
    import pypdf
except ImportError:
    pypdf = None


REFERENCE_HEADERS = [
    'references', 'bibliography', 'works cited', 'sources',
    'المراجع', 'قائمة المراجع', 'المصادر', 'المصادر والمراجع'
]

def is_reference_header(text: str) -> bool:
    clean = re.sub(r'[^\w\s]', '', text).strip().lower()
    return clean in REFERENCE_HEADERS


def extract_from_docx(filepath: str) -> List[Dict[str, Any]]:
    """
    Safely extracts paragraphs and table text from a .docx file without crashing
    on merged cells or complex table layouts.
    Returns list of blocks: [{'text': str, 'type': 'paragraph'|'table_cell', 'is_ref': bool}]
    """
    if docx is None:
        raise RuntimeError("python-docx is not installed.")

    doc = docx.Document(filepath)
    blocks = []
    in_references = False

    # Extract paragraphs
    for p in doc.paragraphs:
        txt = p.text.strip()
        if not txt:
            continue

        if is_reference_header(txt):
            in_references = True

        blocks.append({
            'text': txt,
            'type': 'paragraph',
            'is_ref': in_references
        })

    # Extract tables safely using direct XML tc elements (bypasses python-docx row.cells bugs)
    for tbl in doc.tables:
        for tc in tbl._tbl.xpath('.//w:tc'):
            cell_text = ''.join(tc.itertext()).strip()
            if cell_text:
                blocks.append({
                    'text': cell_text,
                    'type': 'table_cell',
                    'is_ref': in_references
                })

    return blocks


def extract_from_pdf(filepath: str) -> List[Dict[str, Any]]:
    """
    Extracts text blocks from PDF using PyMuPDF (fitz) or PyPDF fallback.
    """
    blocks = []
    in_references = False

    if fitz is not None:
        doc = fitz.open(filepath)
        for page in doc:
            text = page.get_text("text")
            for line in text.split('\n'):
                txt = line.strip()
                if not txt:
                    continue
                if is_reference_header(txt):
                    in_references = True
                blocks.append({
                    'text': txt,
                    'type': 'paragraph',
                    'is_ref': in_references
                })
        doc.close()
    elif pypdf is not None:
        reader = pypdf.PdfReader(filepath)
        for page in reader.pages:
            text = page.extract_text() or ""
            for line in text.split('\n'):
                txt = line.strip()
                if not txt:
                    continue
                if is_reference_header(txt):
                    in_references = True
                blocks.append({
                    'text': txt,
                    'type': 'paragraph',
                    'is_ref': in_references
                })
    else:
        raise RuntimeError("Neither PyMuPDF nor PyPDF is installed.")

    return blocks


def extract_from_txt(filepath: str) -> List[Dict[str, Any]]:
    """
    Extracts text from a plain text file using UTF-8, CP1256 (Arabic), or Latin-1.
    """
    encodings = ['utf-8', 'utf-8-sig', 'cp1256', 'latin-1']
    content = ""
    for enc in encodings:
        try:
            with open(filepath, 'r', encoding=enc) as f:
                content = f.read()
                break
        except UnicodeDecodeError:
            continue

    blocks = []
    in_references = False
    for line in content.split('\n'):
        txt = line.strip()
        if not txt:
            continue
        if is_reference_header(txt):
            in_references = True
        blocks.append({
            'text': txt,
            'type': 'paragraph',
            'is_ref': in_references
        })
    return blocks


def extract_document(filepath: str) -> Dict[str, Any]:
    """
    Main extraction function supporting .docx, .pdf, .txt.
    Returns structured document representation:
    {
        'filepath': str,
        'filename': str,
        'blocks': [...],
        'total_words': int,
        'char_count': int
    }
    """
    ext = os.path.splitext(filepath)[1].lower()
    filename = os.path.basename(filepath)

    if ext == '.docx':
        blocks = extract_from_docx(filepath)
    elif ext == '.pdf':
        blocks = extract_from_pdf(filepath)
    elif ext in ['.txt', '.rtf']:
        blocks = extract_from_txt(filepath)
    else:
        raise ValueError(f"Unsupported file type: {ext}")

    total_words = sum(len(b['text'].split()) for b in blocks)
    char_count = sum(len(b['text']) for b in blocks)

    return {
        'filepath': filepath,
        'filename': filename,
        'blocks': blocks,
        'total_words': total_words,
        'char_count': char_count
    }
