import io
from typing import Dict, Any
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


def generate_excel_report(analysis_result: Dict[str, Any]) -> io.BytesIO:
    """
    Generates a comprehensive, beautifully styled multi-sheet Excel report:
    1. Overview & Ranking (ملخص الأبحاث ونسب التشابه)
    2. Sentences to Rephrase (الجمل المطلوب تعديلها في كل بحث)
    3. Cross-Similarity Matrix (مصفوفة التشابه بين الملفات)
    """
    wb = openpyxl.Workbook()
    # remove default sheet
    wb.remove(wb.active)

    # Styles
    font_title = Font(name='Segoe UI', size=14, bold=True, color='FFFFFF')
    font_header = Font(name='Segoe UI', size=11, bold=True, color='FFFFFF')
    font_regular = Font(name='Segoe UI', size=10)
    font_bold = Font(name='Segoe UI', size=10, bold=True)

    fill_primary = PatternFill(start_color='1E293B', end_color='1E293B', fill_type='solid') # Slate 800
    fill_header = PatternFill(start_color='334155', end_color='334155', fill_type='solid')  # Slate 700

    fill_high_risk = PatternFill(start_color='FEE2E2', end_color='FEE2E2', fill_type='solid') # Red light
    font_high_risk = Font(name='Segoe UI', size=10, bold=True, color='991B1B')

    fill_mod_risk = PatternFill(start_color='FEF3C7', end_color='FEF3C7', fill_type='solid') # Yellow light
    font_mod_risk = Font(name='Segoe UI', size=10, bold=True, color='92400E')

    fill_low_risk = PatternFill(start_color='DCFCE7', end_color='DCFCE7', fill_type='solid') # Green light
    font_low_risk = Font(name='Segoe UI', size=10, bold=True, color='166534')

    border_thin = Border(
        left=Side(style='thin', color='E2E8F0'),
        right=Side(style='thin', color='E2E8F0'),
        top=Side(style='thin', color='E2E8F0'),
        bottom=Side(style='thin', color='E2E8F0')
    )

    # -------------------------------------------------------------
    # SHEET 1: Overview
    # -------------------------------------------------------------
    ws1 = wb.create_sheet(title="الملخص ونسب التشابه")
    ws1.views.sheetView[0].rightToLeft = True

    # Title
    ws1.merge_cells('A1:G1')
    cell_t1 = ws1['A1']
    cell_t1.value = "تقرير فحص التشابه الداخلي بين الأبحاث (Cross-Similarity Report)"
    cell_t1.font = font_title
    cell_t1.fill = fill_primary
    cell_t1.alignment = Alignment(horizontal='center', vertical='center')
    ws1.row_dimensions[1].height = 40

    # Summary Stats row
    summary = analysis_result.get('summary', {})
    stats_text = (
        f"إجمالي الأبحاث: {summary.get('total_documents', 0)} | "
        f"عالية الخطورة (>25%): {summary.get('high_risk_count', 0)} | "
        f"متوسطة الخطورة (15-25%): {summary.get('moderate_risk_count', 0)} | "
        f"آمنة (<15%): {summary.get('safe_count', 0)} | "
        f"متوسط نسبة التشابه: {summary.get('average_similarity', 0)}%"
    )
    ws1.merge_cells('A2:G2')
    cell_t2 = ws1['A2']
    cell_t2.value = stats_text
    cell_t2.font = font_bold
    cell_t2.alignment = Alignment(horizontal='center', vertical='center')
    ws1.row_dimensions[2].height = 25

    headers1 = [
        "اسم البحث / الملف",
        "نسبة Turnitin المتوقعة",
        "تصنيف Turnitin الملون",
        "التشابه الداخلي بين الطلاب %",
        "التشابه مع الإنترنت %",
        "الجمل المقتبسة",
        "الكلمات المشتركة",
        "إجمالي الكلمات",
        "أهم الأبحاث المتشابه معها"
    ]
    for col_idx, h in enumerate(headers1, 1):
        cell = ws1.cell(row=4, column=col_idx, value=h)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = Alignment(horizontal='center', vertical='center')
    ws1.row_dimensions[4].height = 28

    curr_row = 5
    for doc in analysis_result.get('documents', []):
        t_pred = doc.get('turnitin_pred', {})
        t_label = t_pred.get('turnitin_label_ar', doc['risk'])
        t_pct = f"{t_pred.get('overall_similarity_index', doc['similarity_percentage'])}%"
        bd = t_pred.get('breakdown', {})
        student_pct = f"{bd.get('student_papers_pct', doc['similarity_percentage'])}%"
        net_pct = f"{bd.get('internet_sources_pct', 0)}%"

        top_srcs = ", ".join([f"{s['other_filename']} ({s['percentage']}%)" for s in doc['top_sources'][:3]]) or "لا يوجد"

        row_vals = [
            doc['filename'],
            t_pct,
            t_label,
            student_pct,
            net_pct,
            doc['flagged_sentences_count'],
            doc['flagged_words_count'],
            doc['total_words'],
            top_srcs
        ]

        for col_idx, val in enumerate(row_vals, 1):
            cell = ws1.cell(row=curr_row, column=col_idx, value=val)
            cell.font = font_regular
            cell.border = border_thin
            cell.alignment = Alignment(vertical='center')

            # Highlight risk & percentage
            if col_idx in [2, 3]:
                if doc['risk'] == 'high':
                    cell.fill = fill_high_risk
                    cell.font = font_high_risk
                elif doc['risk'] == 'moderate':
                    cell.fill = fill_mod_risk
                    cell.font = font_mod_risk
                else:
                    cell.fill = fill_low_risk
                    cell.font = font_low_risk
                cell.alignment = Alignment(horizontal='center', vertical='center')

        ws1.row_dimensions[curr_row].height = 24
        curr_row += 1

    # -------------------------------------------------------------
    # SHEET 2: Sentences to Rephrase
    # -------------------------------------------------------------
    ws2 = wb.create_sheet(title="الجمل المطلوب إعادة صياغتها")
    ws2.views.sheetView[0].rightToLeft = True

    ws2.merge_cells('A1:F1')
    cell_t_s2 = ws2['A1']
    cell_t_s2.value = "قائمة الجمل المقتبسة والمطلوب إعادة صياغتها لكل بحث"
    cell_t_s2.font = font_title
    cell_t_s2.fill = fill_primary
    cell_t_s2.alignment = Alignment(horizontal='center', vertical='center')
    ws2.row_dimensions[1].height = 40

    headers2 = [
        "اسم البحث",
        "الجملة المقتبسة المطلوب تغييرها",
        "عدد الكلمات",
        "نوع التطابق",
        "نسبة التطابق",
        "الملفات التي وردت فيها هذه الجملة"
    ]
    for col_idx, h in enumerate(headers2, 1):
        cell = ws2.cell(row=3, column=col_idx, value=h)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = Alignment(horizontal='center', vertical='center')
    ws2.row_dimensions[3].height = 28

    curr_row2 = 4
    for doc in analysis_result.get('documents', []):
        for s in doc.get('flagged_sentences', []):
            match_type_ar = "تطابق تام (100%)" if s['best_type'] == 'exact' else f"إعادة صياغة جزئية ({int(s['best_similarity'] * 100)}%)"
            other_files = ", ".join(list(set(m['other_filename'] for m in s['matches'])))

            vals = [
                doc['filename'],
                s['text'],
                s['word_count'],
                match_type_ar,
                f"{int(s['best_similarity'] * 100)}%",
                other_files
            ]

            for col_idx, val in enumerate(vals, 1):
                cell = ws2.cell(row=curr_row2, column=col_idx, value=val)
                cell.font = font_regular
                cell.border = border_thin
                cell.alignment = Alignment(vertical='center', wrap_text=(col_idx == 2))

                if col_idx == 4 and s['best_type'] == 'exact':
                    cell.fill = fill_high_risk
                    cell.font = font_high_risk
                    cell.alignment = Alignment(horizontal='center', vertical='center')
                elif col_idx in [3, 4, 5]:
                    cell.alignment = Alignment(horizontal='center', vertical='center')

            ws2.row_dimensions[curr_row2].height = 30
            curr_row2 += 1

    # -------------------------------------------------------------
    # SHEET 3: Similarity Matrix
    # -------------------------------------------------------------
    ws3 = wb.create_sheet(title="مصفوفة التشابه المتبادل")
    ws3.views.sheetView[0].rightToLeft = False

    matrix_data = analysis_result.get('matrix', {})
    filenames = matrix_data.get('filenames', [])
    data_rows = matrix_data.get('data', [])

    # Header Row
    ws3.cell(row=1, column=1, value="Files / Documents").font = font_header
    ws3.cell(row=1, column=1).fill = fill_primary
    for col_idx, fn in enumerate(filenames, 2):
        cell = ws3.cell(row=1, column=col_idx, value=fn[:25] + '...')
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = Alignment(horizontal='center', vertical='center')

    for row_idx, (fn, row_vals) in enumerate(zip(filenames, data_rows), 2):
        header_cell = ws3.cell(row=row_idx, column=1, value=fn)
        header_cell.font = font_bold
        header_cell.border = border_thin
        for col_idx, val in enumerate(row_vals, 2):
            cell = ws3.cell(row=row_idx, column=col_idx, value=f"{val}%")
            cell.font = font_regular
            cell.border = border_thin
            cell.alignment = Alignment(horizontal='center', vertical='center')

            # Color gradient
            if row_idx - 2 == col_idx - 2:
                # Diagonal (same file)
                cell.fill = PatternFill(start_color='F1F5F9', end_color='F1F5F9', fill_type='solid')
            elif val >= 25.0:
                cell.fill = fill_high_risk
                cell.font = font_high_risk
            elif val >= 15.0:
                cell.fill = fill_mod_risk
                cell.font = font_mod_risk
            elif val > 0:
                cell.fill = fill_low_risk
                cell.font = font_low_risk

    # Auto-adjust column widths
    for ws in [ws1, ws2, ws3]:
        for col in ws.columns:
            col_letter = get_column_letter(col[0].column)
            if ws == ws2 and col_letter == 'B':
                ws.column_dimensions[col_letter].width = 60  # wide for sentence text
            else:
                max_len = max(len(str(cell.value or '')) for cell in col[:20])
                ws.column_dimensions[col_letter].width = max(12, min(max_len + 4, 45))

    stream = io.BytesIO()
    wb.save(stream)
    stream.seek(0)
    return stream
